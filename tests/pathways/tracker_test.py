import importlib.util, tempfile, unittest, json
from pathlib import Path
from openpyxl import Workbook
spec=importlib.util.spec_from_file_location('tracker_extractor',Path(__file__).resolve().parents[2]/'scripts/extract-pathway-tracker.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
class TrackerTests(unittest.TestCase):
 def make(self,root,formula=False):
  w=Workbook();w.remove(w.active)
  for name in module.SUPPORTED:w.create_sheet(name)
  for s in w:s['A2']='Coach: PRIVATE NAME';s['A3']='Client: PRIVATE NAME'
  for name,blocks,n in [(module.SUPPORTED[0],[(4,2,5,11),(4,9,5,11),(13,2,14,20),(13,9,14,20)],6),(module.SUPPORTED[1],[(4,2,5,8),(4,9,5,8),(10,2,11,14),(10,9,11,14)],6),(module.SUPPORTED[2],[(4,2,5,9),(4,8,5,9),(11,2,12,16),(11,8,12,16)],5)]:
   s=w[name]
   for header,col,start,end in blocks:
    for i in range(n):s.cell(header,col+i,'SYNTHETIC HEADER '+str(i))
    for i,row in enumerate(range(start,end+1)):
     s.cell(row,col,'SYNTHETIC ROW '+str(i));s.cell(row,col+1,'=1+1' if formula else 'SYNTHETIC VALUE')
  for s,pos in zip(w,['A22','A16','A18']):s[pos]='SYNTHETIC NOTE'
  p=Path(root)/'input.xlsx';w.save(p);return p
 def test_preserves_cells_and_repeats_without_identifiers(self):
  with tempfile.TemporaryDirectory()as d:
   b=module.extract_tracker(self.make(d));self.assertEqual(len(b['observations']),67);self.assertNotIn('PRIVATE NAME',json.dumps(b));self.assertEqual(len({o['duplicateGroup'] for o in b['observations'] if o['duplicateGroup']}),3)
 def test_formulas_are_source_text_not_executed(self):
  with tempfile.TemporaryDirectory()as d:
   b=module.extract_tracker(self.make(d,True));f=b['observations'][0]['fields'][1];self.assertEqual(f['raw'],'=1+1');self.assertEqual(f['reading'],'uncertain')
 def test_blank_cells_stay_missing(self):
  with tempfile.TemporaryDirectory()as d:
   b=module.extract_tracker(self.make(d));f=b['observations'][0]['fields'][2];self.assertIsNone(f['raw']);self.assertEqual(f['reading'],'blank')
 def test_goal_and_identity_not_inferred(self):
  with tempfile.TemporaryDirectory()as d:
   b=module.extract_tracker(self.make(d));self.assertEqual(b['sources'][0]['goal'],'not_stated');self.assertEqual(b['sources'][0]['subjectAssociation'],'unconfirmed')
if __name__=='__main__':unittest.main()
