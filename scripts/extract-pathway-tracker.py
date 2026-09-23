"""Offline extraction of the supported three-sheet Master Tracker layout.

No model/API calls, no formulas executed, no automatic identity matching.
Output is SENSITIVE even after removing the dedicated name/coach header rows.
Review free text for identifiers before importing. Never commit generated output.
"""
import argparse
import hashlib
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SUPPORTED = ('Daily Passive Tracking', 'Weekly Gym Log', 'Biometrics Baseline')

def extract_tracker(path: Path):
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 10_000_000:
        raise ValueError('Use an XLSX workbook up to 10 MB.')
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > 25_000_000 or len(z.infolist()) > 1000:
            raise ValueError('Workbook archive exceeds extraction limits.')
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    if any(n not in workbook.sheetnames for n in SUPPORTED):
        raise ValueError('Unsupported layout. Do not silently guess field mappings.')
    observations, issues, signatures = [], [], {}
    def field(label, cell):
        value = cell.value
        return {'label': str(label), 'raw': None if value is None else str(value),
                'reading': 'blank' if value is None else ('uncertain' if cell.data_type == 'f' else 'clear'), 'unit': None}
    def add_block(sheet_name, header_row, first_column, start, end, group):
        sheet = workbook[sheet_name]
        headers = [sheet.cell(header_row, c).value for c in range(first_column, first_column+6)]
        if any(v is None for v in headers):
            raise ValueError('Incomplete expected headers; manual layout review required.')
        content_signature = []
        added = []
        for row in range(start, end+1):
            cells = [sheet.cell(row,c) for c in range(first_column,first_column+6)]
            # Values and original coordinates remain separate from interpretations.
            item = {'id': f'tracker-{len(observations)+1:03d}', 'sourceId':'tracker-workbook',
                    'locator':f'{sheet_name}!{cells[0].coordinate}:{cells[-1].coordinate}',
                    'section':sheet_name, 'group':group,
                    'rowLabel':str(cells[0].value) if cells[0].value is not None else '(blank row label)',
                    'fields':[field(headers[i],cell) for i,cell in enumerate(cells)], 'duplicateGroup':None}
            observations.append(item);added.append(item)
            content_signature.append([None if c.value is None else str(c.value) for c in cells])
        # Do not deduplicate. A repeated block can be a template or a real repeated observation.
        sig = (sheet_name, json.dumps(content_signature,ensure_ascii=False))
        if sig in signatures:
            previous, key = signatures[sig]
            for item in previous + added: item['duplicateGroup'] = key
        else:
            signatures[sig] = (added, 'repeat-'+str(len(signatures)+1))
    for index,(header,col,start,end) in enumerate([(4,2,5,11),(4,9,5,11),(13,2,14,20),(13,9,14,20)],1):
        add_block(SUPPORTED[0],header,col,start,end,f'Week {index} / printed block {index}')
    for index,(header,col,start,end) in enumerate([(4,2,5,8),(4,9,5,8),(10,2,11,14),(10,9,11,14)],1):
        add_block(SUPPORTED[1],header,col,start,end,f'Printed block {index}; chronology not established')
    # Biometrics has five columns rather than six.
    sheet = workbook[SUPPORTED[2]]
    bio_signatures = {}
    for index,(header,col,start,end) in enumerate([(4,2,5,9),(4,8,5,9),(11,2,12,16),(11,8,12,16)],1):
        values=[]; added=[]
        for row in range(start,end+1):
            cells=[sheet.cell(row,c) for c in range(col,col+5)]
            item={'id':f'tracker-{len(observations)+1:03d}','sourceId':'tracker-workbook',
                  'locator':f'{SUPPORTED[2]}!{cells[0].coordinate}:{cells[-1].coordinate}',
                  'section':SUPPORTED[2],'group':f'Printed block {index}; chronology not established',
                  'rowLabel':str(cells[0].value),
                  'fields':[field(sheet.cell(header,col+i).value,c) for i,c in enumerate(cells)],'duplicateGroup':None}
            observations.append(item);added.append(item);values.append([None if c.value is None else str(c.value) for c in cells])
        sig=json.dumps(values,ensure_ascii=False)
        if sig in bio_signatures:
            for item in bio_signatures[sig]+added:item['duplicateGroup']='biometrics-repeated-block'
        else:bio_signatures[sig]=added
    for sheet_name,coordinate in zip(SUPPORTED,['A22','A16','A18']):
        cell=workbook[sheet_name][coordinate]
        observations.append({'id':f'tracker-{len(observations)+1:03d}','sourceId':'tracker-workbook',
            'locator':f'{sheet_name}!{coordinate}','section':sheet_name,'group':'Printed source note','rowLabel':'Source note',
            'fields':[field('Note',cell)],'duplicateGroup':None})
    issues = [
        {'id':'tracker-identity','sourceIds':['tracker-workbook'],'code':'subject_not_linked','detail':'This workbook has its own client header. No association with other uploads or Firebase accounts has been confirmed. Dedicated coach/client headers are omitted. Free text still needs a privacy review.','severity':'blocker','status':'open','resolution':None},
        {'id':'tracker-repeats','sourceIds':['tracker-workbook'],'code':'repeated_blocks','detail':'Repeated printed blocks are preserved, not counted as extra sessions, independent progress, or additional customers. Confirm whether populated entries are real observations or examples.','severity':'blocker','status':'open','resolution':None},
        {'id':'tracker-metric','sourceIds':['tracker-workbook'],'code':'metric_definition_missing','detail':'Preserve Sticky Jump, Dead Hang, 45-2 and all original labels. Measurement protocol and the meaning of the jump value are not established. Do not reinterpret the value as jump height.','severity':'warning','status':'open','resolution':None},
        {'id':'tracker-calories','sourceIds':['tracker-workbook'],'code':'observation_not_target','detail':'An estimated calorie note is not a calorie target, dietary prescription, or verified nutrient measurement. Blank cells are missing, not zero.','severity':'warning','status':'open','resolution':None},
        {'id':'tracker-age','sourceIds':['tracker-workbook'],'code':'age_and_goal_not_confirmed','detail':'The workbook does not establish an age, an adult eligibility check, or a fat-loss goal. Language about school must not be converted into an inferred age or a diet recommendation.','severity':'blocker','status':'open','resolution':None}
    ]
    return {'schemaVersion':1,'kind':'private_pathway_source','usage':'source_review_only',
            'privacy':{'directIdentifiersRemoved':True,'containsHealthData':True},
            'sources':[{'id':'tracker-workbook','kind':'tracker','subjectGroup':'subject-tracker-unlinked','subjectAssociation':'unconfirmed','goal':'not_stated','goalBasis':'not_stated'}],
            'observations':observations,'issues':issues}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args()
    if 'training-pathways-private' not in args.out.parts:
        raise SystemExit('Output must be inside a training-pathways-private folder (gitignored).')
    if args.out.exists():raise SystemExit('Output exists. Choose a new file; do not overwrite reviewed source data.')
    bundle=extract_tracker(args.workbook)
    args.out.parent.mkdir(parents=True,exist_ok=True)
    args.out.write_text(json.dumps(bundle,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'sheets':3,'observations':len(bundle['observations']),'output':str(args.out),'databaseWrites':0,'privacyReviewRequired':True}))
if __name__=='__main__':main()
